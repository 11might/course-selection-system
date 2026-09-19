# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

docs_dir = Path(__file__).resolve().parent


def set_run_font(run, size=11, bold=False, mono=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    if mono:
        run.font.name = "Consolas"
        r = run._element
        rPr = r.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:ascii"), "Consolas")
        rFonts.set(qn("w:hAnsi"), "Consolas")
        rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    else:
        run.font.name = "Microsoft YaHei"
        r = run._element
        rPr = r.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_paragraph(doc, text, size=11, bold=False, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    return p


def append_md_to_doc(doc: Document, md_path: Path, section_title: str | None = None):
    if section_title:
        add_paragraph(doc, section_title, size=16, bold=True, space_after=10)
        add_paragraph(doc, "————————", size=10, space_after=8)
    # reuse converter by writing temp logic: call md_to_docx pieces
    text = md_path.read_text(encoding="utf-8")
    # strip first H1 to avoid duplicate giant titles when merging
    lines = text.splitlines()
    if lines and lines[0].startswith("# "):
        lines = lines[1:]
    # run through same loop as md_to_docx — duplicate minimal by writing to a helper
    _feed_lines(doc, lines)


def _feed_lines(doc: Document, lines):
    in_code = False
    code_lines = []

    def flush_code():
        nonlocal code_lines
        if not code_lines:
            return
        block = "\n".join(code_lines)
        p = doc.add_paragraph()
        run = p.add_run(block)
        set_run_font(run, size=9, mono=True)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.left_indent = Pt(12)
        code_lines = []

    for raw in lines:
        line = raw.rstrip("\n")

        if line.strip().startswith("```"):
            if in_code:
                flush_code()
                in_code = False
            else:
                in_code = True
                code_lines = []
            continue

        if in_code:
            code_lines.append(line)
            continue

        if not line.strip():
            continue

        if line.startswith("# "):
            add_paragraph(doc, line[2:].strip().replace("**", ""), size=16, bold=True, space_after=10)
            continue

        if line.startswith("## "):
            add_paragraph(doc, line[3:].strip().replace("**", ""), size=14, bold=True, space_after=8)
            continue

        if line.startswith("### "):
            add_paragraph(doc, line[4:].strip().replace("**", ""), size=12, bold=True, space_after=6)
            continue

        if line.startswith("---"):
            add_paragraph(doc, "————————", size=10, space_after=8)
            continue

        if line.startswith("> "):
            add_paragraph(doc, line[2:].strip().replace("**", ""), size=10, space_after=6)
            continue

        if line.startswith("|"):
            body = line.replace("|", "").replace("-", "").replace(":", "").strip()
            if body == "":
                continue
            cells = [c.strip().replace("**", "") for c in line.strip("|").split("|")]
            add_paragraph(doc, "  |  ".join(cells), size=10, space_after=2)
            continue

        stripped = line.lstrip()
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(stripped[2:].replace("**", ""))
            set_run_font(run, size=11)
            continue

        if len(stripped) >= 3 and stripped[0].isdigit() and ". " in stripped[:4]:
            p = doc.add_paragraph(style="List Number")
            content = stripped.split(". ", 1)[-1].replace("**", "")
            run = p.add_run(content)
            set_run_font(run, size=11)
            continue

        clean = line.replace("**", "")
        add_paragraph(doc, clean, size=11, space_after=6)

    flush_code()


def md_to_docx(md_path: Path, docx_path: Path, title: str):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    add_paragraph(doc, title, size=18, bold=True, space_after=12)
    _feed_lines(doc, md_path.read_text(encoding="utf-8").splitlines())
    doc.save(docx_path)
    print("saved", docx_path)


def main():
    progress_md = docs_dir / "学习进度与技术说明.md"
    journal_md = docs_dir / "今日总结-2026-07-25.md"

    files = [
        (progress_md, docs_dir / "学习进度与技术说明.docx", "选课系统 · 学习进度与技术说明"),
        (
            journal_md,
            docs_dir / "学习日记汇总.docx",
            "选课系统 · 学习日记汇总（按时间续写，含 7/25 起全部）",
        ),
    ]
    for md, docx, title in files:
        md_to_docx(md, docx, title)

    # 合并一份，方便手机只拷一个文件
    combined = Document()
    style = combined.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    add_paragraph(combined, "选课系统 · 学习总结完整版（进度说明 + 日记）", size=18, bold=True, space_after=12)
    append_md_to_doc(combined, progress_md, "第一部分：学习进度与技术说明")
    append_md_to_doc(combined, journal_md, "第二部分：学习日记（按时间续写）")
    combined_path = docs_dir / "学习总结完整版.docx"
    combined.save(combined_path)
    print("saved", combined_path)
    print("done")


if __name__ == "__main__":
    main()
